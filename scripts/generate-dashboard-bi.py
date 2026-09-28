"""
Génère un classeur BI GymAdvisor à partir d'un export
(onglets Ressenti / Charges UA / Charges par zone).
Les graphiques et tableaux s'adaptent au nombre de lignes.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

LIME = "C8F135"
DARK = "121214"
CARD = "1F1F23"
MUTED = "A0A0A8"
FG = "ECECEE"
BLUE = "6EA8FF"
AMBER = "F5A524"
WHITE = "FFFFFF"
LIGHT_BG = "F4F4F5"
LIGHT_CARD = "FFFFFF"
LIGHT_BORDER = "D2D5DB"
LIGHT_MUTED = "5B6470"
LIGHT_FG = "1A1A1D"

thin = Border(
    left=Side(style="thin", color=LIGHT_BORDER),
    right=Side(style="thin", color=LIGHT_BORDER),
    top=Side(style="thin", color=LIGHT_BORDER),
    bottom=Side(style="thin", color=LIGHT_BORDER),
)


def style_header(cell, fill=LIME, font_color=DARK):
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=True, color=font_color, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_kpi_label(cell):
    cell.font = Font(bold=True, color=LIGHT_MUTED, size=9)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_kpi_value(cell):
    cell.font = Font(bold=True, color=LIGHT_FG, size=18)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def autosize(ws, min_width=10, max_width=28):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def read_export(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    ressenti = wb["Ressenti"]
    ua = wb["Charges UA"]
    zones = wb["Charges par zone"]

    summary = {
        "periode": ressenti["A2"].value,
        "energie": ressenti["B2"].value,
        "sommeil": ressenti["C2"].value,
        "douleurs": ressenti["D2"].value,
        "motivation": ressenti["E2"].value,
        "nb_checkins": ressenti["F2"].value,
        "score_ressenti": ressenti["G2"].value,
        "kpi_seances": None,
        "kpi_volume": None,
        "kpi_rpe": None,
    }
    for row in ressenti.iter_rows(min_row=1, max_row=ressenti.max_row, max_col=2, values_only=True):
        if row[0] == "KPI séances":
            summary["kpi_seances"] = row[1]
        elif row[0] == "KPI volume kg":
            summary["kpi_volume"] = row[1]
        elif row[0] == "KPI RPE moyen":
            summary["kpi_rpe"] = row[1]

    feeling_series = []
    header_row = None
    for i, row in enumerate(
        ressenti.iter_rows(min_row=1, max_row=ressenti.max_row, max_col=6, values_only=True),
        start=1,
    ):
        if row[0] == "Bucket":
            header_row = i
            break
    if header_row:
        for row in ressenti.iter_rows(
            min_row=header_row + 1,
            max_row=ressenti.max_row,
            max_col=6,
            values_only=True,
        ):
            if row[0] is None:
                continue
            feeling_series.append(
                {
                    "bucket": row[0],
                    "energie": row[1],
                    "sommeil": row[2],
                    "douleurs": row[3],
                    "motivation": row[4],
                    "nb": row[5] or 0,
                }
            )

    ua_summary = {
        "athlete": ua["B1"].value,
        "total_ua": ua["B2"].value,
        "aigu": ua["B3"].value,
        "chronique": ua["B4"].value,
        "ratio": ua["B5"].value,
    }
    ua_series = []
    ua_header = None
    for i, row in enumerate(ua.iter_rows(min_row=1, max_row=ua.max_row, max_col=3, values_only=True), start=1):
        if row[0] == "Bucket":
            ua_header = i
            break
    if ua_header:
        for row in ua.iter_rows(min_row=ua_header + 1, max_row=ua.max_row, max_col=3, values_only=True):
            if row[0] is None:
                continue
            ua_series.append({"bucket": row[0], "ua": row[1] or 0, "rpe": row[2]})

    zone_rows = []
    for row in zones.iter_rows(min_row=2, max_row=zones.max_row, max_col=3, values_only=True):
        if row[0] is None:
            continue
        zone_rows.append({"zone": row[0], "tonnage": row[1] or 0, "percent": row[2] or 0})

    return {
        "summary": summary,
        "feeling_series": feeling_series,
        "ua_summary": ua_summary,
        "ua_series": ua_series,
        "zones": zone_rows,
    }


def write_data_sheets(wb: Workbook, data: dict):
    # --- Ressenti ---
    ws = wb.create_sheet("Ressenti")
    s = data["summary"]
    headers = [
        "Période",
        "Énergie",
        "Sommeil",
        "Douleurs",
        "Motivation",
        "Nb check-ins",
        "Score ressenti",
    ]
    ws.append(headers)
    for c in range(1, 8):
        style_header(ws.cell(1, c))
    ws.append(
        [
            s["periode"],
            s["energie"],
            s["sommeil"],
            s["douleurs"],
            s["motivation"],
            s["nb_checkins"],
            s["score_ressenti"],
        ]
    )
    ws.append([])
    ws.append(["KPI séances", s["kpi_seances"]])
    ws.append(["KPI volume kg", s["kpi_volume"]])
    ws.append(["KPI RPE moyen", s["kpi_rpe"]])
    ws.append([])
    series_header_row = 8
    ws.append(["Bucket", "Énergie", "Sommeil", "Douleurs", "Motivation", "Nb"])
    for c in range(1, 7):
        style_header(ws.cell(series_header_row, c), fill="1A1A1D", font_color=LIME)
    for item in data["feeling_series"]:
        ws.append(
            [
                item["bucket"],
                item["energie"],
                item["sommeil"],
                item["douleurs"],
                item["motivation"],
                item["nb"],
            ]
        )
    last_feel = series_header_row + len(data["feeling_series"])
    if data["feeling_series"]:
        table = Table(
            displayName="TbRessentiSerie",
            ref=f"A{series_header_row}:F{last_feel}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        ws.add_table(table)
    autosize(ws)

    # --- Charges UA ---
    wu = wb.create_sheet("Charges UA")
    u = data["ua_summary"]
    wu.append(["Athlète", u["athlete"]])
    wu.append(["Total UA", u["total_ua"]])
    wu.append(["Aigu 7j", u["aigu"]])
    wu.append(["Chronique /sem", u["chronique"]])
    wu.append(["Ratio A/C", u["ratio"]])
    wu.append([])
    ua_header_row = 7
    wu.append(["Bucket", "UA", "RPE moyen"])
    for c in range(1, 4):
        style_header(wu.cell(ua_header_row, c), fill="1A1A1D", font_color=LIME)
    for item in data["ua_series"]:
        wu.append([item["bucket"], item["ua"], item["rpe"]])
    last_ua = ua_header_row + len(data["ua_series"])
    if data["ua_series"]:
        table = Table(displayName="TbChargesUA", ref=f"A{ua_header_row}:C{last_ua}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        wu.add_table(table)
        wu.conditional_formatting.add(
            f"B{ua_header_row + 1}:B{last_ua}",
            DataBarRule(
                start_type="num",
                start_value=0,
                end_type="max",
                color=BLUE,
            ),
        )
    autosize(wu)

    # --- Zones ---
    wz = wb.create_sheet("Charges par zone")
    wz.append(["Zone", "Tonnage (kg)", "%"])
    for c in range(1, 4):
        style_header(wz.cell(1, c), fill="1A1A1D", font_color=LIME)
    for item in data["zones"]:
        wz.append([item["zone"], item["tonnage"], item["percent"]])
    last_z = 1 + len(data["zones"])
    if data["zones"]:
        table = Table(displayName="TbZones", ref=f"A1:C{last_z}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True
        )
        wz.add_table(table)
        wz.conditional_formatting.add(
            f"C2:C{last_z}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                end_type="max",
                end_color=LIME,
            ),
        )
    autosize(wz)

    return {
        "feel_header": series_header_row,
        "feel_last": last_feel if data["feeling_series"] else series_header_row,
        "ua_header": ua_header_row,
        "ua_last": last_ua if data["ua_series"] else ua_header_row,
        "zone_last": last_z if data["zones"] else 1,
    }


def build_dashboard(wb: Workbook, data: dict, spans: dict):
    ws = wb.create_sheet("Dashboard", 0)
    s = data["summary"]
    u = data["ua_summary"]

    ws.sheet_view.showGridLines = False
    ws["A1"] = "GymAdvisor — Dashboard BI"
    ws["A1"].font = Font(bold=True, size=20, color=LIGHT_FG)
    ws.merge_cells("A1:H1")

    ws["A2"] = u["athlete"] or "Sportif"
    ws["A2"].font = Font(size=14, color=LIGHT_MUTED)
    ws.merge_cells("A2:D2")

    ws["E2"] = s["periode"] or ""
    ws["E2"].font = Font(size=12, color=LIGHT_MUTED)
    ws["E2"].alignment = Alignment(horizontal="right")
    ws.merge_cells("E2:H2")

    # KPI cards row 4-5
    kpis = [
        ("A4", "B4", "Séances", s["kpi_seances"] if s["kpi_seances"] is not None else "—"),
        ("C4", "D4", "Volume (kg)", s["kpi_volume"] if s["kpi_volume"] is not None else "—"),
        ("E4", "F4", "RPE moyen", s["kpi_rpe"] if s["kpi_rpe"] is not None else "—"),
        (
            "G4",
            "H4",
            "Score ressenti",
            s["score_ressenti"] if s["score_ressenti"] is not None else "—",
        ),
    ]
    for label_cell, value_cell, label, value in kpis:
        ws[label_cell] = label
        style_kpi_label(ws[label_cell])
        ws[label_cell].fill = PatternFill("solid", fgColor="EEF0F3")
        ws[value_cell] = value
        style_kpi_value(ws[value_cell])
        ws[value_cell].fill = PatternFill("solid", fgColor="EEF0F3")
        ws[label_cell].border = thin
        ws[value_cell].border = thin

    ws["A6"] = "Total UA"
    ws["B6"] = u["total_ua"]
    ws["C6"] = "Aigu 7 j"
    ws["D6"] = u["aigu"]
    ws["E6"] = "Chronique / sem"
    ws["F6"] = u["chronique"]
    ws["G6"] = "Ratio A/C"
    ws["H6"] = u["ratio"]
    for col in range(1, 9):
        cell = ws.cell(6, col)
        cell.font = Font(bold=col % 2 == 1, color=LIGHT_MUTED if col % 2 == 1 else LIGHT_FG, size=10 if col % 2 == 1 else 14)
        cell.border = thin
        cell.fill = PatternFill("solid", fgColor="F7F8FA")

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 28
    ws.row_dimensions[6].height = 24

    # Charts — ressenti (line) from TbRessentiSerie sheet
    wr = wb["Ressenti"]
    feel_start = spans["feel_header"]
    feel_end = spans["feel_last"]
    if feel_end > feel_start:
        line = LineChart()
        line.title = "Ressenti pré-séance (série)"
        line.style = 10
        line.y_axis.title = "Score /5"
        line.x_axis.title = None
        line.height = 10
        line.width = 15
        cats = Reference(wr, min_col=1, min_row=feel_start + 1, max_row=feel_end)
        data_ref = Reference(wr, min_col=2, max_col=5, min_row=feel_start, max_row=feel_end)
        line.add_data(data_ref, titles_from_data=True)
        line.set_categories(cats)
        line.shape = 4
        ws.add_chart(line, "A8")

    # UA bar + RPE line combo-ish: bar for UA
    wu = wb["Charges UA"]
    ua_start = spans["ua_header"]
    ua_end = spans["ua_last"]
    if ua_end > ua_start:
        bar = BarChart()
        bar.type = "col"
        bar.title = "Charges UA par période"
        bar.y_axis.title = "UA"
        bar.style = 10
        bar.height = 10
        bar.width = 12
        data_ua = Reference(wu, min_col=2, min_row=ua_start, max_row=ua_end)
        cats_ua = Reference(wu, min_col=1, min_row=ua_start + 1, max_row=ua_end)
        bar.add_data(data_ua, titles_from_data=True)
        bar.set_categories(cats_ua)

        line_rpe = LineChart()
        line_rpe.y_axis.axId = 200
        line_rpe.y_axis.title = "RPE"
        data_rpe = Reference(wu, min_col=3, min_row=ua_start, max_row=ua_end)
        line_rpe.add_data(data_rpe, titles_from_data=True)
        line_rpe.set_categories(cats_ua)
        line_rpe.y_axis.crosses = "max"
        bar.y_axis.crosses = "min"
        bar += line_rpe
        ws.add_chart(bar, "I8")

    # Zones pie
    wz = wb["Charges par zone"]
    zone_last = spans["zone_last"]
    if zone_last > 1:
        pie = PieChart()
        pie.title = "Répartition tonnage par zone"
        labels = Reference(wz, min_col=1, min_row=2, max_row=zone_last)
        data_z = Reference(wz, min_col=2, min_row=1, max_row=zone_last)
        pie.add_data(data_z, titles_from_data=True)
        pie.set_categories(labels)
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.dataLabels.showVal = False
        pie.dataLabels.showCatName = False
        pie.height = 10
        pie.width = 12
        ws.add_chart(pie, "A24")

        bar_z = BarChart()
        bar_z.type = "bar"
        bar_z.title = "Tonnage par zone (kg)"
        bar_z.style = 10
        bar_z.height = 10
        bar_z.width = 12
        data_zb = Reference(wz, min_col=2, min_row=1, max_row=zone_last)
        cats_z = Reference(wz, min_col=1, min_row=2, max_row=zone_last)
        bar_z.add_data(data_zb, titles_from_data=True)
        bar_z.set_categories(cats_z)
        ws.add_chart(bar_z, "I24")

    ws["A40"] = (
        "Source : onglets Ressenti, Charges UA, Charges par zone (même format que l’export GymAdvisor). "
        "Les tableaux Excel (TbRessentiSerie, TbChargesUA, TbZones) s’étendent si tu ajoutes des lignes "
        "sous les en-têtes Bucket / Zone. Les graphiques du Dashboard sont recalculés à chaque génération."
    )
    ws["A40"].font = Font(size=9, color=LIGHT_MUTED, italic=True)
    ws.merge_cells("A40:H42")
    ws["A40"].alignment = Alignment(wrap_text=True, vertical="top")

    for letter, width in {
        "A": 14,
        "B": 14,
        "C": 14,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 12,
    }.items():
        ws.column_dimensions[letter].width = width


def write_guide(wb: Workbook):
    ws = wb.create_sheet("Mode d'emploi", 1)
    ws["A1"] = "Comment utiliser ce Dashboard BI"
    ws["A1"].font = Font(bold=True, size=16)
    lines = [
        "",
        "1. Cet onglet Dashboard est généré automatiquement à partir des 3 onglets de données.",
        "2. Format attendu (identique à l’export coach GymAdvisor) :",
        "   - Ressenti : résumé période + KPI + tableau série (Bucket, Énergie, Sommeil, Douleurs, Motivation, Nb)",
        "   - Charges UA : athlète, totaux, ratio A/C + série (Bucket, UA, RPE moyen)",
        "   - Charges par zone : Zone, Tonnage (kg), %",
        "3. Volume de données : les graphiques et tableaux Excel s’adaptent au nombre de lignes présentes.",
        "4. Pour mettre à jour : regénère l’export depuis GymAdvisor, ou relance le script generate-dashboard-bi.py",
        "   en passant le chemin du nouveau fichier .xlsx.",
        "5. Tu peux aussi coller de nouvelles lignes dans les tableaux Tb* (Excel étend le tableau).",
        "",
        "Astuce : ouvre d’abord l’onglet Dashboard pour le suivi, puis les onglets de détail pour forer.",
    ]
    for i, line in enumerate(lines, start=2):
        ws[f"A{i}"] = line
        ws[f"A{i}"].font = Font(size=11, color=LIGHT_FG)
    ws.column_dimensions["A"].width = 110


def build(source: Path, dest: Path):
    data = read_export(source)
    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    spans = write_data_sheets(wb, data)
    build_dashboard(wb, data, spans)
    write_guide(wb)

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print(f"OK -> {dest}")


if __name__ == "__main__":
    downloads = Path.home() / "Downloads"
    default_src = downloads / "suivi-lilian-month-2026-09-01 (1).xlsx"
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else default_src
    if not src.exists():
        raise SystemExit(f"Fichier source introuvable: {src}")

    repo = Path(__file__).resolve().parents[1]
    out_repo = repo / "templates" / "GymAdvisor-Dashboard-BI.xlsx"
    out_dl = downloads / "GymAdvisor-Dashboard-BI.xlsx"
    build(src, out_repo)
    build(src, out_dl)
